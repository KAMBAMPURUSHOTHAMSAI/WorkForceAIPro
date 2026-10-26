import csv
import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any

from fastapi import Response

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape


def export_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, Decimal):
        return str(value)

    if isinstance(value, (dict, list)):
        return json.dumps(value, default=str)

    return str(value)


def export_report(
    title: str,
    filename: str,
    columns: list[tuple[str, str]],
    rows: list[dict[str, Any]],
    export_format: str,
) -> Response:
    export_format = export_format.lower()

    headers = [label for _, label in columns]

    data = [
        [
            export_value(row.get(field))
            for field, _ in columns
        ]
        for row in rows
    ]

    if export_format == "csv":
        buffer = StringIO(newline="")
        writer = csv.writer(buffer)

        writer.writerow(headers)
        writer.writerows(data)

        return Response(
            content=("\ufeff" + buffer.getvalue()).encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}.csv"'
            },
        )

    if export_format == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Report"

        worksheet.append(headers)

        for row in data:
            worksheet.append(row)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="17365D",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column_index, (_, label) in enumerate(columns, start=1):
            column_values = [
                len(str(worksheet.cell(row=row_index, column=column_index).value or ""))
                for row_index in range(1, worksheet.max_row + 1)
            ]

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(max(column_values, default=10) + 2, 12), 35)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return Response(
            content=output.getvalue(),
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}.xlsx"'
            },
        )

    if export_format == "pdf":
        output = BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            leftMargin=8 * mm,
            rightMargin=8 * mm,
            topMargin=10 * mm,
            bottomMargin=10 * mm,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            fontSize=16,
            alignment=TA_LEFT,
            spaceAfter=8,
        )

        cell_style = ParagraphStyle(
            "ReportCell",
            parent=styles["BodyText"],
            fontSize=6,
            leading=8,
            wordWrap="CJK",
        )

        header_style = ParagraphStyle(
            "ReportHeader",
            parent=cell_style,
            textColor=colors.white,
            fontName="Helvetica-Bold",
        )

        table_data = [
            [
                Paragraph(escape(value), header_style)
                for value in headers
            ]
        ]

        for row in data:
            table_data.append(
                [
                    Paragraph(escape(value), cell_style)
                    for value in row
                ]
            )

        available_width = landscape(A4)[0] - 16 * mm
        column_width = available_width / max(len(headers), 1)

        table = Table(
            table_data,
            colWidths=[column_width] * len(headers),
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )

        document.build(
            [
                Paragraph(escape(title), title_style),
                Paragraph(
                    f"Total records: {len(rows)}",
                    styles["Normal"],
                ),
                Spacer(1, 8),
                table,
            ]
        )

        output.seek(0)

        return Response(
            content=output.getvalue(),
            media_type="application/pdf",
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}.pdf"'
            },
        )

    raise ValueError("Unsupported export format")